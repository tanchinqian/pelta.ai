from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\USER\Desktop\Hack-Attack\pelta-ai-demonstration-script.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles

def set_font(style, size, color="000000", bold=False):
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold

normal = styles["Normal"]
set_font(normal, 11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, before, after, color in [
    ("Heading 1", 20, 20, 6, "000000"),
    ("Heading 2", 16, 18, 6, "000000"),
    ("Heading 3", 14, 16, 4, "434343"),
]:
    st = styles[name]
    set_font(st, size, color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ["List Bullet", "List Number"]:
    st = styles[name]
    set_font(st, 11)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
r = title.add_run("pelta.ai Demonstration Script")
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(0, 0, 0)
r.font.bold = False

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
run = sub.add_run("Total runtime: 2 minutes 30 seconds | Employee and administrator views")
run.font.name = "Arial"
run.font.size = Pt(11)
run.font.color.rgb = RGBColor.from_string("555555")

doc.add_heading("Pre-demo setup", level=1)
for text in [
    "Open the pelta.ai landing page and sign in to ChatGPT with the Prompt Guard extension enabled.",
    "Keep the Classify Tool, Access Requests, and Dashboard pages ready for quick navigation.",
    "Use the sample prompt and Otter.ai classification details below to keep the demonstration consistent.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Part 1 - Employee View: Protection and Redress", level=1)
p = doc.add_paragraph()
p.add_run("Timing: ").bold = True
p.add_run("0:00-1:10")

steps = [
    ("0:00-0:10 | Introduce the employee view", "On the pelta.ai landing page, click Employee.",
     "We will begin from the employee's perspective. Pelta.ai enables staff to use generative AI safely without requiring them to understand every security policy."),
    ("0:10-0:35 | Demonstrate the active interceptor", "Open ChatGPT. Paste the sample prompt below, then click ChatGPT's Send button.",
     "When the employee clicks Send, Pelta.ai intercepts the prompt before it reaches the AI provider. The first DLP pass scans for deterministic patterns such as emails, phone numbers, credentials, and financial identifiers."),
    ("0:35-0:50 | Explain the verdict", "Point to the blocked overlay, highlighted text, risk level, and detection method.",
     "The prompt is blocked because it contains identifiable contact information. The employee can see exactly what triggered the decision, its risk level, and whether it was detected by regex or contextual LLM analysis."),
    ("0:50-1:10 | Show explanation and redress", "Return to pelta.ai, click Right to Explanation, and select the relevant blocked event. Click Suggest Safe Alternatives, then show Submit Appeal.",
     "The employee is not left at a dead end. They can review the decision, generate safer versions of the prompt, or submit an appeal through the EU AI Act redress workflow. Every appeal stays linked to the original security event."),
]

for heading, cue, speech in steps:
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph()
    p.add_run("CLICK CUE: ").bold = True
    p.add_run(cue)
    p = doc.add_paragraph()
    p.add_run("SAY: ").bold = True
    p.add_run(speech)

doc.add_heading("Sample blocked prompt", level=2)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.25)
p.paragraph_format.right_indent = Inches(0.25)
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Draft a handover email to sarah.lee@company.com. Her phone number is +1-415-555-0192, and the internal project budget is confidential.")
r.font.name = "Arial"
r.font.size = Pt(10.5)
r.font.italic = True

doc.add_page_break()

doc.add_heading("Part 2 - Admin View: Classification, Approval, and Audit", level=1)
p = doc.add_paragraph()
p.add_run("Timing: ").bold = True
p.add_run("1:10-2:30")

steps2 = [
    ("1:10-1:18 | Switch perspectives", "Return to the landing page and click Admin.",
     "We will now switch to the governance administrator's view, where tool risk, approvals, security events, and redress requests are managed centrally."),
    ("1:18-1:45 | Classify a new AI tool", "Click Classify Tool. Enter Otter.ai and the use case below, then click Classify Tool. Point to the risk tier, data categories, NIST functions, justification, and recommended policy.",
     "Instead of manually researching every new AI service, the administrator submits its intended business use. Pelta.ai retrieves relevant NIST AI RMF context through its lightweight RAG layer and sends that grounding to Gemini."),
    ("1:45-2:00 | Explore the Tool Registry", "Click Tool Registry. Search for Otter.ai, then open its detail panel. Point to its approval status, risk tier, NIST mapping, data categories, justification, and policy. Briefly show the Edit or Reclassify action.",
     "The classification is immediately added to the central Tool Registry. Administrators can search and filter the catalog, inspect why each tool received its risk rating, override its policy, change its approval status, or rerun the Gemini assessment as its use case evolves."),
    ("2:00-2:15 | Review approvals and redress", "Click Access Requests, open Redress Appeals, and select a pending appeal. Click Approve, or briefly show the rejection-comment field.",
     "Administrators review AI-tool requests and employee appeals with the original context, detected categories, business justification, and decision history. An approval or explained rejection is automatically added to the audit trail."),
    ("2:15-2:25 | Show governance analytics", "Click Dashboard. Point to verdict distribution, risk charts, NIST coverage, and the Tool Registry summary.",
     "The dashboard consolidates prompt verdicts, risk distribution, detection methods, NIST coverage, request activity, and the live tool registry. All visualizations use the same operational records created by these workflows."),
    ("2:25-2:30 | Close", "Return focus to the dashboard overview.",
     "Together, these two views give employees safe access to AI while giving administrators complete, explainable governance and oversight."),
]

for heading, cue, speech in steps2:
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph()
    p.add_run("CLICK CUE: ").bold = True
    p.add_run(cue)
    p = doc.add_paragraph()
    p.add_run("SAY: ").bold = True
    p.add_run(speech)

doc.add_heading("Tool classification input", level=2)
for label, value in [
    ("Tool name", "Otter.ai"),
    ("Description", "AI meeting transcription and summarization for internal leadership and customer calls"),
]:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)

doc.core_properties.title = "pelta.ai Demonstration Script"
doc.core_properties.subject = "Two-minute-thirty-second employee and administrator product demonstration"
doc.core_properties.author = "pelta.ai"
doc.save(OUT)
print(OUT)
